from io import BytesIO
from decimal import Decimal, InvalidOperation
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def _d(value, places: str = "0.01") -> str:
    if value is None or value == "":
        return "N/A"
    try:
        dec = value if isinstance(value, Decimal) else Decimal(str(value))
    except InvalidOperation:
        return str(value)
    quant = Decimal(places)
    return str(dec.quantize(quant)) if dec == dec else str(dec)

def _dsigned(value) -> str:
    if value is None or value == "":
        return "N/A"
    try:
        dec = value if isinstance(value, Decimal) else Decimal(str(value))
    except InvalidOperation:
        return str(value)
    formatted = _d(dec)
    return f"+{formatted}" if dec >= 0 else formatted

def _set_cell_background(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def _set_table_borders(table, color="cbd5e1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def generate_docx_certificate(report_data: Dict[str, Any], observations: List[Dict[str, Any]]) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    is_approved = str(report_data.get("status", "")).upper() == "APPROVED"
    
    # Header
    head_p = doc.add_paragraph()
    head_p.paragraph_format.space_after = Pt(4)
    run_auth = head_p.add_run("LEGAL METROLOGY DIRECTORIES\n")
    run_auth.bold = True
    run_auth.font.size = Pt(13)
    run_auth.font.color.rgb = RGBColor(15, 23, 42)
    run_sub = head_p.add_run("OIML R-76 Technical Compliance Authority")
    run_sub.font.size = Pt(9)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc_title = "VERIFICATION CERTIFICATE OF CONFORMITY" if is_approved else f"TEST REPORT ({report_data.get('status', 'DRAFT')})"
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(8)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run(doc_title)
    t_run.bold = True
    t_run.font.size = Pt(15)

    ref_p = doc.add_paragraph()
    ref_p.paragraph_format.space_after = Pt(14)
    ref_run = ref_p.add_run(f"Certificate No: CERT-{report_data['id']:05d} | Standard: OIML R-76-1 | Status: {report_data.get('status', 'DRAFT')}")
    ref_run.font.size = Pt(9)
    ref_run.font.color.rgb = RGBColor(71, 85, 105)

    # Specifications
    doc.add_paragraph().add_run("1. INSTRUMENT & VERIFICATION SPECIFICATIONS").bold = True
    specs = [
        ("Manufacturer", str(report_data.get("manufacturer", "N/A")), "Accuracy Class", f"Class {report_data.get('accuracy_class', 'N/A')}"),
        ("Model Number", str(report_data.get("model", "N/A")), "Max Capacity", f"{_d(report_data.get('capacity'))} {report_data.get('unit', '')}"),
        ("Serial Number", str(report_data.get("serial_number", "N/A")), "Scale Interval (d)", f"{_d(report_data.get('d'))} {report_data.get('unit', '')}"),
        ("Verification Interval (e)", f"{_d(report_data.get('e'))} {report_data.get('unit', '')}", "Resolution (n)", f"{_d(report_data.get('n'), '1')} e"),
        ("Test Location", str(report_data.get("test_location") or "N/A"), "Reference Eqp", str(report_data.get("reference_equipment") or "N/A")),
    ]
    table_specs = doc.add_table(rows=len(specs), cols=4)
    table_specs.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table_specs)

    for i, row in enumerate(specs):
        for col_idx, text in enumerate(row):
            cell = table_specs.cell(i, col_idx)
            cell.text = text
            if col_idx in (0, 2):
                cell.paragraphs[0].runs[0].bold = True
                _set_cell_background(cell, "F8FAFC")
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Observations
    doc.add_paragraph().add_run("2. EVALUATION & METROLOGICAL OBSERVATIONS").bold = True
    headers = ["Test Category", "Load Position", "Ref Load", "Indicated Value", "Error (E)", "Allowed MPE", "Result"]
    unit = report_data.get("unit", "")
    table_obs = doc.add_table(rows=1 + len(observations), cols=7)
    _set_table_borders(table_obs)

    for col_idx, heading in enumerate(headers):
        cell = table_obs.cell(0, col_idx)
        cell.text = heading
        _set_cell_background(cell, "E2E8F0")
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)

    for row_idx, o in enumerate(observations, start=1):
        row_data = [
            str(o.get("test_type", "")),
            str(o.get("position") or "Centre"),
            f"{_d(o.get('reference_load'))} {unit}",
            f"{_d(o.get('observed_value'))} {unit}",
            _dsigned(o.get("calculated_error")),
            f"±{_d(o.get('mpe'))}",
            str(o.get("status", "")),
        ]
        for col_idx, val in enumerate(row_data):
            cell = table_obs.cell(row_idx, col_idx)
            cell.text = val
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            if col_idx == 6:
                cell.paragraphs[0].runs[0].bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Summary
    doc.add_paragraph().add_run("3. COMPLIANCE DETERMINATION & SIGN-OFF").bold = True
    summary_rows = [
        ("Overall Metrological Status", str(report_data.get("overall_result", "PENDING"))),
        ("Evaluating Technician", str(report_data.get("technician_name", "N/A"))),
        ("Approval Authority / Signatory", str(report_data.get("approved_by") or "Pending Approval")),
        ("Rule Version Applied", str(report_data.get("rule_version", "N/A"))),
    ]
    table_sum = doc.add_table(rows=len(summary_rows), cols=2)
    _set_table_borders(table_sum)

    for i, (k, v) in enumerate(summary_rows):
        c0 = table_sum.cell(i, 0)
        c1 = table_sum.cell(i, 1)
        c0.text = k
        c1.text = v
        _set_cell_background(c0, "F8FAFC")
        c0.paragraphs[0].runs[0].bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(8.5)
        c1.paragraphs[0].runs[0].font.size = Pt(8.5)

    buf = BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()
    buf.close()
    return docx_bytes